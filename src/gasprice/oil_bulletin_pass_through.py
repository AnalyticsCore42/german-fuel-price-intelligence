"""
Module: oil_bulletin_pass_through.py
Description: Tankrabatt pass-through for the dashboard tiles, difference-in-differences on the
EU Weekly Oil Bulletin (net-of-tax pump prices, Mondays).

  outcome  = German net price minus the mean of euro-area control countries without an excise/VAT
             change, price regulation or pump rebate around the event (cost shocks cancel out)
  estimate = 1 - [outcome, first 2 weeks in force - outcome, 2 weeks before start] / net energy-tax cut

Placebo RMSE (same windows, 2015+ weeks without a German tax event, in pp of a 14.04 ct cut):
Super 95 8.6 calm / 12.3 turbulent, diesel 9.6 / 14.7; the raw pump change scores 12.2-29.4.
"""

from __future__ import annotations

from datetime import date, datetime, timedelta

import numpy as np
import polars as pl
import requests

from gasprice.pass_through_model import DATA_DIR, TAX_CUT_EVENTS, VAT_GROSS_UP, energy_tax_cut

URL = ("https://energy.ec.europa.eu/document/download/906e60ca-8b6a-44e7-8589-652854d2fd3f_en"
       "?filename=Weekly_Oil_Bulletin_Prices_History_maticni_4web.xlsx")
RAW_FILE = DATA_DIR / "eu_oil_bulletin" / "Weekly_Oil_Bulletin_Prices_History.xlsx"
OUT_FILE = DATA_DIR / "dashboard" / "tankrabatt_pass_through.parquet"

EURO_AREA = ["AT", "BE", "CY", "EE", "ES", "FI", "FR", "GR", "HR", "IE", "IT", "LT", "LU", "LV",
             "MT", "NL", "PT", "SI", "SK"]
REGULATED = {"BE", "LU", "SI", "HR", "MT", "CY"}
REBATES = {"FR": (date(2022, 4, 1), date(2022, 12, 31)), "ES": (date(2022, 4, 1), date(2022, 12, 31))}
# product key in the workbook -> (column in the excise/VAT sheets, tax type, dashboard fuels)
PRODUCTS = {"euro95": (2, "petrol", ("e5", "e10")), "diesel": (3, "diesel", ("diesel",))}
WINDOW = timedelta(days=14)


def download(path=RAW_FILE) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    resp = requests.get(URL, timeout=120)
    resp.raise_for_status()
    path.write_bytes(resp.content)


def _prices(wb, product: str) -> dict[date, dict[str, float]]:
    rows = list(wb["Prices wo taxes"].iter_rows(values_only=True))
    cols = {h.split("_")[0]: i for i, h in enumerate(rows[0])
            if isinstance(h, str) and h.endswith(f"price_wo_tax_{product}")}
    return {r[0].date(): {cc: r[i] / 1000 if isinstance(r[i], (int, float)) else np.nan for cc, i in cols.items()}
            for r in rows[3:] if isinstance(r[0], datetime)}


def _change_dates(wb, sheet: str, col: int) -> dict[str, list[date]]:
    hist, cc = {}, None
    for r in wb[sheet].iter_rows(min_row=5, values_only=True):
        if isinstance(r[0], str) and r[0].endswith("_"):
            cc = r[0][:-1]
        if cc and isinstance(r[1], datetime):
            hist.setdefault(cc, []).append((r[1].date(), r[col]))
    return {cc: [d for (d, v), (_, v0) in zip(sorted(h)[1:], sorted(h)[:-1]) if v != v0] for cc, h in hist.items()}


def _controls(lo: date, hi: date, changes: dict[str, list[date]]) -> list[str]:
    return [cc for cc in EURO_AREA
            if cc not in REGULATED
            and not (cc in REBATES and REBATES[cc][0] <= hi and REBATES[cc][1] >= lo)
            and not any(lo <= d <= hi for d in changes.get(cc, []))]


def _mean(px, ccs, lo, hi) -> float:
    vals = [px[d]["DE"] - np.nanmean([px[d].get(c, np.nan) for c in ccs]) for d in px if lo <= d <= hi]
    return float(np.nanmean(vals)) if vals else np.nan


def estimate_all(path=RAW_FILE) -> pl.DataFrame:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = []
    for product, (col, tax_type, fuels) in PRODUCTS.items():
        px = _prices(wb, product)
        excise, vat = _change_dates(wb, "Excise duties", col), _change_dates(wb, "VAT", col)
        changes = {cc: excise.get(cc, []) + vat.get(cc, []) for cc in set(excise) | set(vat)}
        last = max(px)
        for ev in TAX_CUT_EVENTS:
            if ev.start + timedelta(days=7) > last:
                continue
            ccs = _controls(ev.announced - timedelta(weeks=8), ev.end + timedelta(weeks=4), changes)
            cut = energy_tax_cut(ev, tax_type)
            shift = _mean(px, ccs, ev.start, ev.start + WINDOW - timedelta(days=1)) - _mean(px, ccs, ev.start - WINDOW, ev.start - timedelta(days=1))
            pct = (1 - shift / cut) * 100
            statutory_ct = cut * VAT_GROSS_UP * 100
            for fuel in fuels:
                rows.append({"event": ev.name, "fuel": fuel, "pass_through_pct": pct,
                             "statutory_ct": statutory_ct, "passed_ct": statutory_ct * pct / 100,
                             "controls": " ".join(ccs), "data_until": last})
    return pl.DataFrame(rows)


def main() -> None:
    try:
        download()
    except requests.RequestException as e:
        # weekly data: last week's file still gives the same window estimates
        if not RAW_FILE.exists():
            raise
        print(f"Oil Bulletin download failed ({e}); using saved file from "
              f"{datetime.fromtimestamp(RAW_FILE.stat().st_mtime):%Y-%m-%d}")
    df = estimate_all()
    OUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    df.write_parquet(OUT_FILE)
    print(df.select("event", "fuel", "pass_through_pct", "passed_ct", "statutory_ct"))


if __name__ == "__main__":
    main()
